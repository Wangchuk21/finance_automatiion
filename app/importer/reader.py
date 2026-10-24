"""Excel file reader and parser for the importer."""

import hashlib
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from app.importer.column_map import (
    parse_month_column,
    get_canonical_column,
    parse_amount,
    parse_date,
    clean_cid,
    clean_mobile,
    reshape_month_columns,
)
from app.importer.validation import ValidationReport, validate_excel_import


@dataclass
class ParsedSheet:
    """Result of parsing an Excel sheet."""
    sheet_name: str
    headers: List[str]
    rows: List[Dict[str, Any]]
    month_columns: Dict[str, Tuple[str, int, int]]  # original_name -> (type, month, year)
    static_columns: Dict[str, str]  # original_name -> canonical_name
    totals_row: Optional[int] = None
    sheet_totals: Dict[str, float] = None


def detect_header_row(worksheet, max_rows: int = 10) -> int:
    """Detect the header row by looking for common column names."""
    header_keywords = {"reference", "ref", "cid", "customer", "mobile", "salary", "external", "received"}
    
    for row_idx in range(1, min(max_rows + 1, worksheet.max_row + 1)):
        row_values = []
        for col_idx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            val = str(cell.value).strip().lower() if cell.value else ""
            row_values.append(val)
        
        # Count how many header keywords match
        matches = sum(1 for v in row_values if any(kw in v for kw in header_keywords))
        if matches >= 3:  # At least 3 recognizable headers
            return row_idx
    
    return 1  # Default to first row


def read_excel_sheet(
    file_path: str,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = None,
) -> ParsedSheet:
    """
    Read and parse an Excel sheet.
    
    Args:
        file_path: Path to the Excel file
        sheet_name: Specific sheet name (None = first sheet)
        header_row: 1-indexed header row (None = auto-detect)
    
    Returns:
        ParsedSheet with parsed data and column mappings
    """
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    
    try:
        if sheet_name:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        
        # Detect header row if not specified
        if header_row is None:
            header_row = detect_header_row(ws)
        
        # Read headers
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            header = str(cell.value).strip() if cell.value else f"col_{col_idx}"
            headers.append(header)
        
        # Map columns
        month_columns = {}
        static_columns = {}
        
        for col_idx, header in enumerate(headers):
            # Try to parse as month column
            month_info = parse_month_column(header)
            if month_info:
                month_columns[header] = month_info
            else:
                # Try to map to canonical static column
                canonical = get_canonical_column(header)
                if canonical:
                    static_columns[header] = canonical
        
        # Read data rows
        rows = []
        sheet_totals = {}
        totals_row = None
        
        for row_idx in range(header_row + 1, ws.max_row + 1):
            row_data = {}
            is_empty = True
            
            for col_idx, header in enumerate(headers):
                cell = ws.cell(row=row_idx, column=col_idx + 1)
                value = cell.value
                if value is not None and str(value).strip():
                    is_empty = False
                row_data[header] = value
            
            if is_empty:
                # Check if this might be a totals row
                first_col = str(row_data.get(headers[0], "")).strip().lower() if headers else ""
                if any(kw in first_col for kw in ["total", "sum", "grand total", "subtotal"]):
                    totals_row = row_idx
                    # Parse totals
                    for col_idx, header in enumerate(headers):
                        cell = ws.cell(row=row_idx, column=col_idx + 1)
                        if cell.value is not None:
                            parsed = parse_amount(cell.value)
                            if parsed is not None:
                                sheet_totals[header] = parsed
                    continue
                # Skip completely empty rows
                continue
            
            rows.append(row_data)
        
        return ParsedSheet(
            sheet_name=ws.title,
            headers=headers,
            rows=rows,
            month_columns=month_columns,
            static_columns=static_columns,
            totals_row=totals_row,
            sheet_totals=sheet_totals,
        )
    
    finally:
        wb.close()


def parse_sheet_to_contracts(parsed: ParsedSheet) -> List[Dict[str, Any]]:
    """
    Convert parsed sheet rows to contract records with monthly installments.
    
    Returns list of contract dicts, each containing:
    - static fields (reference, customer, etc.)
    - installments: list of monthly installment dicts
    """
    contracts = []
    
    for row in parsed.rows:
        # Skip total/summary rows
        ref = str(row.get("Reference", "")).strip().lower()
        if ref in ("total", "sum", "grand total", "subtotal"):
            continue
        
        # Build static contract data
        contract = {}
        
        # Map static columns
        for orig_col, canonical in parsed.static_columns.items():
            value = row.get(orig_col)
            
            if canonical == "cid":
                contract[canonical] = clean_cid(value)
            elif canonical == "mobile":
                mobile, _ = clean_mobile(value)
                contract[canonical] = mobile
            elif canonical in ("contract_date",):
                contract[canonical] = parse_date(value)
            elif canonical in ("monthly_deduction", "sales_amount", "gst", "emi_months"):
                parsed_val = parse_amount(value)
                if parsed_val is not None:
                    contract[canonical] = parsed_val
            else:
                contract[canonical] = str(value).strip() if value is not None else None
        
        # Reshape monthly columns
        installments = reshape_month_columns(row, parsed.month_columns)
        
        if not contract.get("reference"):
            continue  # Skip rows without reference
        
        contract["installments"] = installments
        contracts.append(contract)
    
    return contracts


def calculate_file_checksum(file_path: str) -> str:
    """Calculate SHA256 checksum of a file."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


def import_excel_file(
    file_path: str,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = None,
) -> Tuple[List[Dict[str, Any]], ValidationReport, Dict[str, Any]]:
    """
    Complete import pipeline: read, parse, validate.
    
    Returns:
        (contracts, validation_report, preview_summary)
    """
    # Read and parse
    parsed = read_excel_sheet(file_path, sheet_name, header_row)
    
    # Validate using original parsed rows (with original column names)
    report = validate_excel_import(
        parsed.rows,
        parsed.month_columns,
        parsed.static_columns,
        parsed.sheet_totals,
    )
    
    # Convert to contract format
    contracts = parse_sheet_to_contracts(parsed)
    
    # Generate preview using contracts (which have installments)
    from app.importer.validation import generate_preview_summary
    preview = generate_preview_summary(contracts, parsed.month_columns)
    
    return contracts, report, preview